import json
import os
import re
from collections import Counter, OrderedDict
from datetime import datetime, timezone

import openpyxl

ROOT = r'D:\漫画工具整体代码'
SOURCE = r'D:\漫画\万虫11-21话\二审\开局神戒_制作交付包_011-021_V005\开局神戒_制作总表_011-021_V005.xlsx'
OUT = os.path.join(ROOT, 'data', 'asset-registry.json')
REPORT = os.path.join(ROOT, 'data', 'asset-validation.json')


def text(value):
    return '' if value is None else str(value).strip()


def variant_id(note):
    match = re.search(r'(V-[0-9a-f]{16})', note or '')
    return match.group(1) if match else ''


def first_panel(note):
    match = re.search(r'首用\s*([A-Z]\d{3}-P\d{3})', note or '')
    return match.group(1) if match else ''


wb = openpyxl.load_workbook(SOURCE, read_only=True, data_only=True)

# 1. characters
characters = OrderedDict()
for row in list(wb['04_角色资产'].iter_rows(values_only=True))[1:]:
    asset_id = text(row[0])
    if not asset_id:
        continue
    characters[asset_id] = {
        'assetId': asset_id,
        'type': 'character',
        'name': text(row[1]),
        'aliases': text(row[2]),
        'firstAppearance': text(row[3]),
        'firstNamed': text(row[4]),
        'ageBody': text(row[5]),
        'faceLock': text(row[6]),
        'hair': text(row[7]),
        'baseCostumeId': text(row[8]),
        'signatureProp': text(row[9]),
        'expression': text(row[10]),
        'reference': text(row[11]),
        'prompt': text(row[12]),
        'consistency': text(row[13]),
        'note': text(row[14]),
        'specStatus': 'pending',
        'stage': 'backlog',
        'approval': 'candidate',
        'costumes': [],
        'variants': [],
        'components': {},
        'bindings': []
    }

# 2. costumes
costumes = OrderedDict()
for row in list(wb['05_服装形态'].iter_rows(values_only=True))[1:]:
    costume_id = text(row[0])
    role_id = text(row[1])
    if not costume_id or not role_id:
        continue
    key = f'{role_id}-{costume_id}'
    costumes[key] = {
        'assetId': key,
        'type': 'costume',
        'characterId': role_id,
        'costumeId': costume_id,
        'name': text(row[2]),
        'stage': text(row[3]),
        'material': text(row[4]),
        'mainColor': text(row[5]),
        'accentColor': text(row[6]),
        'pattern': text(row[7]),
        'states': {
            '完整': text(row[8]),
            '破损': text(row[9]),
            '血迹': text(row[10]),
            '湿身': text(row[11]),
            '变身': text(row[12])
        },
        'reference': text(row[13]),
        'prompt': text(row[14]),
        'note': text(row[15]),
        'specStatus': 'pending',
        'stage': 'backlog',
        'approval': 'candidate',
        'components': {},
        'bindings': []
    }
    if role_id in characters:
        characters[role_id]['costumes'].append(key)

# 3. scenes
scenes = OrderedDict()
for row in list(wb['06_场景资产'].iter_rows(values_only=True))[1:]:
    scene_id = text(row[0])
    if not scene_id:
        continue
    if scene_id not in scenes:
        scenes[scene_id] = {
            'assetId': scene_id,
            'type': 'scene',
            'name': text(row[1]),
            'specStatus': 'pending',
            'stage': 'backlog',
            'approval': 'candidate',
            'components': {
                'floorPlan': {'status': 'missing'},
                'master': {'status': 'missing'},
                'details': {'status': 'missing'}
            },
            'variants': [],
            'bindings': []
        }
    scenes[scene_id]['variants'].append({
        'camera': text(row[2]),
        'time': text(row[3]),
        'weather': text(row[4]),
        'light': text(row[5]),
        'architecture': text(row[6]),
        'foreground': text(row[7]),
        'midground': text(row[8]),
        'background': text(row[9]),
        'reuse': text(row[10]),
        'reference': text(row[11]),
        'prompt': text(row[12]),
        'note': text(row[13]),
        'variantId': variant_id(text(row[13])),
        'firstPanel': first_panel(text(row[13])),
        'status': 'pending_design'
    })

# 4. props
props = OrderedDict()
for row in list(wb['07_道具资产'].iter_rows(values_only=True))[1:]:
    prop_id = text(row[0])
    name = text(row[1])
    if not prop_id:
        continue
    key = f'{prop_id}|{name}' if name else prop_id
    if key not in props:
        owner = ''
        match = re.search(r'[（(]([^）)]+)[）)]', name)
        if match:
            owner = match.group(1)
        props[key] = {
            'assetId': prop_id,
            'type': 'prop',
            'name': name,
            'owner': owner,
            'kind': text(row[2]),
            'material': text(row[3]),
            'states': {
                '完整': text(row[4]),
                '破损': text(row[5]),
                '激活': text(row[6]),
                '隐藏': text(row[7])
            },
            'holdRule': text(row[8]),
            'quantity': text(row[9]),
            'ledgerId': text(row[10]),
            'reference': text(row[11]),
            'prompt': text(row[12]),
            'note': text(row[13]),
            'specStatus': 'pending',
            'stage': 'backlog',
            'approval': 'candidate',
            'components': {
                'main': {'status': 'missing'},
                'scabbard': {'status': 'missing'},
                'sheathed': {'status': 'missing'},
                'active': {'status': 'missing'}
            },
            'variants': [],
            'bindings': []
        }
    props[key]['variants'].append({
        'prompt': text(row[12]),
        'note': text(row[13]),
        'variantId': variant_id(text(row[13])),
        'firstPanel': first_panel(text(row[13])),
        'status': 'pending_design'
    })

# 5. known assets
known_characters = {
    'CH003': {
        'identityCard': ('candidate', 'asset_c41a6235-55f2-4a00-a56a-489db91398b0', '6dbebeee-6a6f-4032-a553-429b5651b439'),
        'hands': ('candidate', 'asset_2a3a3ce6-dce3-413e-a1cd-766ba1e711e7', 'bcc31e93-695d-4bd5-980f-debf7b9ed52a')
    },
    'CH002': {
        'identityCard': ('candidate', 'asset_aa5a7563-27dc-4f0d-948d-4d6e9ce02859', '')
    },
    'CH001': {
        'identityCard': ('candidate', 'asset_aed638b1-6b27-46ef-a0c3-4c4d64a7da2a', '')
    }
}
for role_id, components in known_characters.items():
    if role_id not in characters:
        continue
    for name, (status, editor_id, moku_id) in components.items():
        characters[role_id]['components'][name] = {
            'status': status,
            'editorId': editor_id,
            'mokuId': moku_id
        }

known_costumes = {
    'CH001-CO003': {
        'standardCard': ('candidate', 'asset_aed638b1-6b27-46ef-a0c3-4c4d64a7da2a', '')
    },
    'CH002-CO002': {
        'standardCard': ('candidate', 'asset_aa5a7563-27dc-4f0d-948d-4d6e9ce02859', '')
    },
    'CH003-CO001': {
        'standardCard': ('candidate', 'asset_c41a6235-55f2-4a00-a56a-489db91398b0', '6dbebeee-6a6f-4032-a553-429b5651b439')
    },
    'CH003-CO005': {
        'standardCard': ('candidate', 'asset_b76f0ccf-6f81-43e1-b28b-894221695cf4', '7d11670a-51a5-4e04-81a7-924c755300ac'),
        'costumeDesign': ('candidate', 'asset_ef89229e-0148-42bc-aa43-66cd46683528', '3062001d-be1e-43d9-8e40-635bbbf50e54'),
        'hands': ('candidate', 'asset_497ca050-67c9-4a1a-a930-cbfbcfcc15a7', '88cdc58d-1100-4630-b2aa-12e4222234ca')
    }
}
for key, components in known_costumes.items():
    if key not in costumes:
        continue
    for name, (status, editor_id, moku_id) in components.items():
        costumes[key]['components'][name] = {
            'status': status,
            'editorId': editor_id,
            'mokuId': moku_id
        }

# local component files that are not registered in the editor yet
if 'CH003' in characters:
    characters['CH003']['components']['faceCloseup'] = {
        'status': 'candidate',
        'file': 'data/asset-library/standard/face.png'
    }
    characters['CH003']['components']['threeView'] = {
        'status': 'candidate',
        'file': 'data/asset-library/standard/threeview.png'
    }
if 'CH003-CO005' in costumes:
    costumes['CH003-CO005']['components'].update({
        'faceCloseup': {'status': 'candidate', 'file': 'data/asset-library/standard/face.png'},
        'front': {'status': 'candidate', 'file': 'data/asset-library/standard/front.png'},
        'side': {'status': 'candidate', 'file': 'data/asset-library/standard/side.png'},
        'back': {'status': 'candidate', 'file': 'data/asset-library/standard/back.png'}
    })

for key, prop in props.items():
    if prop['assetId'] == 'PR004' and prop['name'].startswith('蓝虹剑'):
        prop['specStatus'] = 'candidate'
        prop['components']['main'] = {
            'status': 'candidate',
            'editorId': 'asset_4918c22e-994c-4051-ab63-b3a315ce4cc7',
            'mokuId': 'ed037b27-f145-411f-bf0c-ab7de8632433'
        }
        prop['components']['scabbard'] = {'status': 'missing'}
        prop['components']['sheathed'] = {'status': 'missing'}

# 6. prompt variants from 09
prompt_rows = list(wb['09_资产提示词'].iter_rows(values_only=True))[1:]
for row in prompt_rows:
    asset_id = text(row[0])
    name = text(row[1])
    kind = text(row[2])
    variant = text(row[3])
    usage = text(row[4])
    prompt = text(row[5])
    negative = text(row[6])
    reference = text(row[7])
    lock = text(row[8])
    note = text(row[9])
    record = {
        'variantId': variant,
        'usage': usage,
        'prompt': prompt,
        'negative': negative,
        'reference': reference,
        'lock': lock,
        'note': note,
        'firstPanel': first_panel(note)
    }
    if kind == 'character' and asset_id in characters:
        characters[asset_id]['variants'].append(record)
    elif kind == 'costume':
        for key, costume in costumes.items():
            if costume['costumeId'] == asset_id or key.endswith(asset_id):
                costume.setdefault('variants', []).append(record)
                break
    elif kind == 'prop':
        for key, prop in props.items():
            if prop['assetId'] == asset_id and prop['name'] == name:
                prop['variants'].append(record)
                break
    elif kind == 'scene' and asset_id in scenes:
        scenes[asset_id]['variants'].append(record)

assets = list(characters.values()) + list(costumes.values()) + list(scenes.values()) + list(props.values())

in_progress = {
    'CH001', 'CH002', 'CH003',
    'CH003-CO001', 'CH003-CO005',
    'PR004', 'SC003'
}
for asset in assets:
    if asset['assetId'] in in_progress or asset.get('costumeId') in ('CO001', 'CO005'):
        asset['stage'] = 'in_progress'
counts = {
    'characters': len(characters),
    'costumes': len(costumes),
    'scenes': len(scenes),
    'props': len(props),
    'assets': len(assets),
    'variants': sum(len(item.get('variants', [])) for item in assets)
}
registry = {
    'schemaVersion': '1.0',
    'source': SOURCE,
    'generatedAt': datetime.now(timezone.utc).isoformat(),
    'counts': counts,
    'assets': assets
}

os.makedirs(os.path.dirname(OUT), exist_ok=True)
with open(OUT, 'w', encoding='utf-8') as handle:
    json.dump(registry, handle, ensure_ascii=False, indent=2)

issues = []


def issue(asset, code, message, severity='error'):
    issues.append({
        'assetId': asset.get('assetId'),
        'type': asset.get('type'),
        'name': asset.get('name'),
        'severity': severity,
        'code': code,
        'message': message
    })


for asset in assets:
    if asset.get('stage') != 'in_progress':
        continue
    kind = asset['type']
    if asset.get('specStatus') in ('pending', ''):
        issue(asset, 'spec_missing', '规格未确认')
    if asset.get('approval') != 'approved':
        issue(asset, 'not_approved', '资产未批准', 'warning')
    if not asset.get('variants'):
        issue(asset, 'variants_missing', '没有 V005 变体记录', 'warning')
    if not asset.get('bindings'):
        issue(asset, 'unbound', '没有画格绑定', 'warning')

    if kind == 'character':
        required = ['identityCard', 'faceCloseup', 'threeView', 'hands']
    elif kind == 'costume':
        required = ['standardCard']
    elif kind == 'scene':
        required = ['floorPlan', 'master', 'details']
    else:
        required = ['main']
        if asset.get('kind') == '武器':
            required += ['scabbard', 'sheathed']

    components = asset.get('components') or {}
    for name in required:
        item = components.get(name)
        if not item or item.get('status') == 'missing':
            issue(asset, f'missing_component:{name}', f'缺少组件 {name}')

summary = {
    'errors': sum(1 for item in issues if item['severity'] == 'error'),
    'warnings': sum(1 for item in issues if item['severity'] == 'warning'),
    'issues': len(issues),
    'byCode': dict(Counter(item['code'] for item in issues).most_common()),
    'byType': dict(Counter(item['type'] for item in issues).most_common()),
    'byAsset': dict(Counter(item['assetId'] for item in issues).most_common(30)),
    'stage': dict(Counter(asset.get('stage', 'backlog') for asset in assets))
}
report = {
    'generatedAt': registry['generatedAt'],
    'summary': summary,
    'counts': counts,
    'issues': issues
}
with open(REPORT, 'w', encoding='utf-8') as handle:
    json.dump(report, handle, ensure_ascii=False, indent=2)

print(json.dumps(counts, ensure_ascii=False))
print(json.dumps(summary, ensure_ascii=False))
print('registry', OUT)
print('report', REPORT)
