"""Read original scripts without modifying the source. No shot splitting."""
import hashlib
import io
import json
import re
from openpyxl import load_workbook


def normalize_script(value, source, index=1):
    shots = value.get('shots')
    if not isinstance(shots, list) or not shots:
        raise ValueError('JSON需要非空shots数组；每格包含code、description，可带key、shotSize、cameraAngle等脚本字段')
    result = []
    episode = str(value.get('episode') or index)
    for i, shot in enumerate(shots, 1):
        if not isinstance(shot, dict) or not str(shot.get('code', '')).strip():
            raise ValueError(f'第{i}格缺少原镜号code')
        row = dict(shot)
        row['key'] = str(shot.get('key') or f'E{episode}-P{i:03}')
        row['code'] = str(shot['code'])
        result.append(row)
    if len({s['key'] for s in result}) != len(result) or len({s['code'] for s in result}) != len(result):
        raise ValueError('镜号code和key都需要唯一，不允许隐式合并')
    return {'episode': episode, 'title': str(value.get('title') or f'第{episode}话'), 'source': source, 'shots': result}


def import_scripts(raw, filename):
    sha = hashlib.sha256(raw).hexdigest()
    base = {'filename': filename, 'sha256': sha}
    if filename.lower().endswith('.json'):
        value = json.loads(raw.decode('utf-8-sig'))
        values = value if isinstance(value, list) else [value]
        return [normalize_script(v, dict(base, kind='json', item=i), i) for i, v in enumerate(values, 1)]
    if not filename.lower().endswith('.xlsx'):
        raise ValueError('当前支持XLSX分镜表或JSON脚本')
    workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    scripts = []
    fields = ['code','characters','location','weather','time','shotSize','cameraAngle','description','dialogue','sfx','note']
    try:
        for sheet in workbook:
            rows = list(sheet.iter_rows(values_only=True))
            header = next((i for i,r in enumerate(rows[:10]) if len(r)>=11 and str(r[0] or '').strip() in ('格号','镜号') and str(r[7] or '').strip() in ('画面描述','画面内容')), None)
            if header is None:
                continue
            aliases = [('格号','镜号'),('出场角色','出现角色'),('当前地点','当前场景'),('当前天气',),('当前时间',),('景别',),('机位角度',),('画面描述','画面内容'),('台词',),('音效字','特效字'),('画面备注',)]
            if any(str(rows[header][j] or '').strip() not in options for j,options in enumerate(aliases)):
                raise ValueError(f'{sheet.title}表头列序不受支持，请按示例整理，不能猜测列对应关系')
            shots = []
            for row_index, row in enumerate(rows[header+1:], header+2):
                if not row or row[0] is None or str(row[0]).strip()=='':
                    continue
                code = str(row[0]).strip()
                if isinstance(row[0], float) and row[0].is_integer():
                    code = str(int(row[0]))
                if not re.fullmatch(r'[\dA-Za-z一二三四五六七八九十零]+(?:[-_.][\dA-Za-z]+)*', code):
                    raise ValueError(f'{sheet.title}!A{row_index}镜号无法可靠识别，请先整理表格')
                item = {k: (str(row[j]) if j<len(row) and row[j] is not None else '') for j,k in enumerate(fields)}
                item.update(code=code, sourceRow=row_index)
                shots.append(item)
            if shots:
                match = re.search(r'\d+', sheet.title)
                episode = match.group() if match else sheet.title
                title = str(rows[0][0] or sheet.title)
                scripts.append(normalize_script({'episode':episode, 'title':title, 'shots':shots}, dict(base,kind='xlsx',sheet=sheet.title,range=f'A{header+2}:K{max(s["sourceRow"] for s in shots)}')))
    finally:
        workbook.close()
    if not scripts:
        raise ValueError('未发现支持的分镜表。表头需按格号、出场角色、当前地点、当前天气、当前时间、景别、机位角度、画面描述、台词、音效字、画面备注排列')
    return scripts
